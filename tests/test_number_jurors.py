"""Tests for skills/number-jurors/scripts/number_jurors.py.

Fixtures are synthetic (no real participants) but copy the quirks seen in
real de novo files: Nelson "Amys Tab" sign-in sheets, ***EB / late / note
suffixes, quoted nicknames, hyphenated surnames, first-name-only answers,
misspellings, duplicate and partial SurveyMonkey submissions.
"""
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills/number-jurors/scripts"))
import number_jurors as nj  # noqa: E402

HEAD = ["Respondent ID", "Collector ID", "Start Date", "End Date", "IP Address",
        "Email Address", "First Name", "Last Name", "Custom Data 1", "rq_flag",
        "What is your name? (First AND Last Name)", "What is your gender?",
        "What is your current age?",
        "What are the greatest sources of stress in your life? (SELECT ALL THAT APPLY.)",
        None, None, "Did you have any problems answering any of the questions?"]
SUB = [None] * 10 + ["Open-Ended Response", "Response", "Open-Ended Response",
                     "Money", "Work", "Other (please specify)", "Open-Ended Response"]

# (SurveyMonkey name, gender, age) in export order (newest first)
SM = [
    ("Ava Thompson", "Female", 31),
    ("Ben Carter", "Male", 44),
    ("chloe diaz ", "Female", 27),
    ("Daniel Evans", "Male", 52),
    ("Grace", "Female", 38),
    ("Henry kellumen", "Male", 61),
    ("Isaac J Moore", "Male", 47),
    ("Julia Nguyen", "Female", 29),
    ("Rachel Adams", "Female", 35),
    ("Rachel Stone", "Female", 41),   # not seated - same first name as J10
    ("Laura Price", "Female", 58),    # absent
    ("Daniel Evans", None, None),     # duplicate, partial
]

# (Attendance, J#, sign-in name)
AMY = [
    ("Present", 1, "Ava Thompson-Reyes"),
    ("Present", 2, 'Benjamin "Ben" Carter ***EB'),
    ("Present", 3, "Chloe Diaz late"),
    ("Present", 4, "Daniel Evans - released at lunch/asleep"),
    ("Present", 5, "Grace Lee"),
    ("Present", 6, "Henry Kellum"),
    ("Present", 7, "Isaac Moore, Jr."),
    ("Present", "8 MIA-CUT", "Julia Nguyen"),
    ("Present", 9, "Kevin Ortiz"),          # never filled out the background
    ("Present", 10, "Rachel Adams"),
    ("ABSENT", None, "Laura Price"),
    ("Released", None, "Rachel Stone"),
]


def make_background(path, head=HEAD, sub=SUB, rows=SM):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(head)
    ws.append(sub)
    ws.merge_cells("N1:P1")
    for i, (nm, g, age) in enumerate(rows):
        rid = 115186000000 + i
        vals = [rid, 440835290, f"2026-09-30 1{i % 10}:00:00", None, f"10.0.0.{i}",
                None, None, None, None, 0.0, nm, g, age,
                "Money" if g else None, None, None, "No" if g else None]
        vals = vals[:len(head)]
        if len(head) > len(vals):
            vals += [None] * (len(head) - len(vals))
        ws.append(vals)
    ws.append([115185999999, 440835290, "2026-09-22 20:11:42", None, "10.9.9.9"])  # partial
    ws.append([115185999998, 440835290, None, None, None, None, None, None, None,
               "Column1", "Column2", "Column3"])  # table artefact
    wb.save(path)


def make_signin(path, amy=AMY, formulas=True):
    wb = openpyxl.Workbook()
    main = wb.active
    main.title = "101 715 AM"
    main.append([])
    main.append([None] * 6 + ["Thursday October 1st"])
    main.append([])
    main.append(["Parking", "Notes", "Q and C", "#", "Notes", "Incentive", "Name", "Gender"])
    for i, (_, _, nm) in enumerate(amy, start=1):
        main.append([None, None, "QC", 30 - i, None, "$425", nm, None])  # '#' is NOT J#
    tab = wb.create_sheet("Amys Tab")
    tab.append(["Attendance", "J#", "Rec Names", "Match/No Match", "Notes",
                None, None, None, "SM Names"])
    for i, (att, num, nm) in enumerate(amy):
        name_cell = f"='101 715 AM'!G{5 + i}" if formulas else nm
        tab.append([att, num, name_cell, None, None, None, None, None, None])
    wb.save(path)


def rows_of(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    return ws, [[c.value for c in r] for r in ws.iter_rows(min_row=3)]


# ------------------------------------------------------------------ units

def test_parse_num():
    assert nj.parse_num(7) == (7, "")
    assert nj.parse_num(7.0) == (7, "")
    assert nj.parse_num("12") == (12, "")
    assert nj.parse_num("J12") == (12, "")
    assert nj.parse_num("21 MIA-CUT") == (21, "MIA-CUT")
    assert nj.parse_num("8 gone too long w tech issues") == (8, "gone too long w tech issues")
    assert nj.parse_num("GHOSTED") == (None, "")
    assert nj.parse_num(None) == (None, "")
    assert nj.parse_num(0) == (None, "")


@pytest.mark.parametrize("raw,first", [
    ('Stephan "Stephanie" Rodriguez ***EB', ["stephan", "rodriguez"]),
    ("Cameron Guthrie late", ["cameron", "guthrie"]),
    ("Morgan Swanson - released at lunch/appeared to be on drugs", ["morgan", "swanson"]),
    ("Richard Edgar, Jr.", ["richard", "edgar"]),
    ("  BRIDGETT RANKIN ", ["bridgett", "rankin"]),
    ("José Peña", ["jose", "pena"]),
])
def test_name_variants_cleaning(raw, first):
    assert nj.name_variants(raw)[0] == first


def test_name_variants_nicknames():
    assert ["stephanie", "rodriguez"] in nj.name_variants('Stephan "Stephanie" Rodriguez')
    assert ["ron", "grindeim"] in nj.name_variants("Ronald \"Ron'' Grindeim")
    assert ["rob", "woodmansee"] in nj.name_variants('Robert " Rob" Woodmansee')


def test_header_regex():
    for h in ["j#", "J #", "juror #", "Juror Number", "Juror Numbers", "juror no.", "Jur #"]:
        assert nj.JUROR_NUM_HEADER.match(nj.norm_header(h)), h
    for h in ["#", "rq_flag", "Juror Name", "Number"]:
        assert not nj.JUROR_NUM_HEADER.match(nj.norm_header(h)), h


# ------------------------------------------------------------------ sign-in

def test_signin_uses_amys_tab_not_recruit_number(tmp_path):
    p = tmp_path / "signin.xlsx"
    make_signin(p)
    seats, where = nj.read_signin(str(p))
    assert where == "Amys Tab!B"
    assert [s.num for s in seats] == [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]
    # names came through the un-calculated ='101 715 AM'!G5 formulas
    assert seats[0].name == "Ava Thompson-Reyes"
    assert seats[7].note == "MIA-CUT"


def test_amys_tab_name_variants(tmp_path):
    for title in ["Amy's Tab", "Amy’s tab", "AMYS TAB", "Amys_Tab"]:
        assert nj._is_amys_tab(title), title
    assert not nj._is_amys_tab("Amys Tab old copy")


def test_only_amys_tab_is_used(tmp_path):
    """A J# column on any other sheet is ignored; no Amy's tab is an error."""
    p = tmp_path / "signin.xlsx"
    make_signin(p)
    wb = openpyxl.load_workbook(p)
    wb["Amys Tab"].title = "Juror Profiles"
    wb.save(p)
    with pytest.raises(SystemExit, match="No Amy's tab"):
        nj.read_signin(str(p))


def test_amys_tab_without_numbers_yet(tmp_path):
    p = tmp_path / "signin.xlsx"
    make_signin(p, amy=[("Present", None, "Ava Thompson")])
    with pytest.raises(SystemExit, match="check-in may not be finished"):
        nj.read_signin(str(p))


# ------------------------------------------------------------------ end to end

def test_blocks_when_a_seated_juror_is_missing(tmp_path, capsys):
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    make_background(bg)
    make_signin(si)
    code = nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out)])
    text = capsys.readouterr().out
    assert code == 2
    assert not out.exists()
    assert "J9 Kevin Ortiz" in text and "NEEDS CONFIRMING" in text


def test_full_run(tmp_path, capsys):
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    make_background(bg)
    make_signin(si)
    code = nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out),
                    "--allow-missing", "9"])
    text = capsys.readouterr().out
    assert code == 0, text
    ws, rows = rows_of(out)
    assert ws.cell(1, 10).value == "Juror Number"
    assert ws.cell(1, 11).value.startswith("What is your name")
    got = [(r[9], r[10].strip()) for r in rows]
    assert got == [(1, "Ava Thompson"), (2, "Ben Carter"), (3, "chloe diaz"),
                   (4, "Daniel Evans"), (5, "Grace"), (6, "Henry kellumen"),
                   (7, "Isaac J Moore"), (8, "Julia Nguyen"), (10, "Rachel Adams")]
    # whole rows move together: Respondent ID still belongs to the same person
    src = {r[10].strip(): r[0] for r in rows_of(bg)[1] if r[10] and r[11]}
    for r in rows:
        assert r[0] == src[r[10].strip()]
    # merged question header survives
    assert "N1:P1" in [str(m) for m in ws.merged_cells.ranges]
    # removed: not seated, duplicate, partial, artefact
    assert "Rachel Stone" in text and "Laura Price" in text
    assert "duplicate submission (1)" in text
    assert 'J8 Julia Nguyen: J# cell also says "MIA-CUT"' in text
    # near matches are flagged for a look
    assert "Henry kellumen" in text and "<- check" in text


def test_two_rachels_never_cross(tmp_path):
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    make_background(bg, rows=[("Rachel Wiggins", "Female", 30), ("Rachel Coleman", "Female", 33)])
    make_signin(si, amy=[("Present", 1, "Rachel Coleman")])
    assert nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out)]) == 0
    _, rows = rows_of(out)
    assert [(r[9], r[10]) for r in rows] == [(1, "Rachel Coleman")]


def test_first_name_only_needs_unique_first_name(tmp_path, capsys):
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    make_background(bg, rows=[("Grace", "Female", 30), ("Grace Kim", "Female", 40)])
    make_signin(si, amy=[("Present", 1, "Grace Lee"), ("Present", 2, "Grace Kim")])
    assert nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out)]) == 2
    assert "J1 Grace Lee" in capsys.readouterr().out
    # once a person confirms it, --map writes it
    assert nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out),
                    "--map", "1=Grace"]) == 0
    _, rows = rows_of(out)
    assert [(r[9], r[10]) for r in rows] == [(1, "Grace"), (2, "Grace Kim")]


def test_existing_juror_number_column_wins_over_rq_flag(tmp_path):
    # 1394 layout: "Juror Number" in column I, rq_flag still in J
    head = HEAD[:8] + ["Juror Number"] + HEAD[9:]
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    make_background(bg, head=head, rows=[("Ava Thompson", "Female", 31)])
    make_signin(si, amy=[("Present", 4, "Ava Thompson")])
    assert nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out)]) == 0
    ws, rows = rows_of(out)
    assert ws.cell(1, 9).value == "Juror Number" and ws.cell(1, 10).value == "rq_flag"
    assert rows[0][8] == 4


def test_inserts_column_when_export_has_none(tmp_path):
    head = [h for h in HEAD if h != "rq_flag"]
    sub = SUB[:9] + SUB[10:]
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(head)
    ws.append(sub)
    ws.merge_cells("M1:O1")
    ws.append([1, 2, None, None, None, None, None, None, None, "Ava Thompson", "Female", 31,
               "Money", None, None, "No"])
    wb.save(bg)
    make_signin(si, amy=[("Present", 3, "Ava Thompson")])
    assert nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out)]) == 0
    ws, rows = rows_of(out)
    assert ws.cell(1, 9).value == "Custom Data 1"
    assert ws.cell(1, 10).value == "Juror Number"
    assert rows[0][9] == 3 and rows[0][10] == "Ava Thompson"
    assert "N1:P1" in [str(m) for m in ws.merged_cells.ranges]


def test_refuses_to_overwrite_input(tmp_path):
    bg, si = tmp_path / "bg.xlsx", tmp_path / "si.xlsx"
    make_background(bg)
    make_signin(si)
    with pytest.raises(SystemExit):
        nj.main(["--background", str(bg), "--signin", str(si), "--out", str(bg)])


def test_duplicate_juror_numbers_rejected(tmp_path):
    bg, si = tmp_path / "bg.xlsx", tmp_path / "si.xlsx"
    make_background(bg)
    make_signin(si, amy=[("Present", 1, "Ava Thompson"), ("Present", 1, "Ben Carter")])
    with pytest.raises(SystemExit, match="used twice"):
        nj.main(["--background", str(bg), "--signin", str(si), "--out", str(tmp_path / "o.xlsx")])


def test_exclude_a_cut_juror(tmp_path, capsys):
    bg, si, out = tmp_path / "bg.xlsx", tmp_path / "si.xlsx", tmp_path / "out.xlsx"
    make_background(bg)
    make_signin(si)
    assert nj.main(["--background", str(bg), "--signin", str(si), "--out", str(out),
                    "--allow-missing", "9", "--exclude", "8"]) == 0
    text = capsys.readouterr().out
    _, rows = rows_of(out)
    assert 8 not in [r[9] for r in rows] and "Julia Nguyen" not in [r[10] for r in rows]
    assert "Left out on request (--exclude): J8 Julia Nguyen" in text
